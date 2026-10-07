from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt


SOURCE = Path(r"D:\Code\知光正版\zhiguang_be-main\tmp\resume-work\resume-source.docx")
OUTPUT = Path(r"D:\Code\知光正版\zhiguang_be-main\output\简历1-项目1深度优化.docx")
GITHUB_URL = "https://github.com/G-Pegasus/zhiguang_be"


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, east_asia, size_pt=10, bold=False, color="000000", underline=False):
    run.bold = bold
    run.underline = underline
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.color.rgb = None
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), east_asia)
    color_el = r_pr.find(qn("w:color"))
    if color_el is None:
        color_el = OxmlElement("w:color")
        r_pr.append(color_el)
    color_el.set(qn("w:val"), color)


def add_text(paragraph, text, *, east_asia="仿宋", size_pt=10, bold=False):
    run = paragraph.add_run(text)
    set_run_font(run, east_asia=east_asia, size_pt=size_pt, bold=bold)
    return run


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "黑体")
    r_pr.append(fonts)

    bold = OxmlElement("w:b")
    r_pr.append(bold)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "18")
    r_pr.append(size_cs)

    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def replace_labeled_paragraph(paragraph, label, body):
    clear_paragraph(paragraph)
    add_text(paragraph, label, east_asia="黑体", size_pt=10, bold=True)
    add_text(paragraph, body, east_asia="仿宋", size_pt=10)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    project_table = document.tables[2]

    title_paragraph = project_table.rows[0].cells[0].paragraphs[0]
    clear_paragraph(title_paragraph)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(
        title_paragraph,
        "知光｜事件驱动知识社区",
        east_asia="黑体",
        size_pt=10,
        bold=True,
    )

    github_paragraph = project_table.rows[0].cells[1].paragraphs[0]
    clear_paragraph(github_paragraph)
    github_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(github_paragraph, "GitHub 145+ Stars", GITHUB_URL)

    background = (
        "面向知识分享场景，构建集内容发布、互动关系、Feed、全文检索与单篇 RAG 问答于一体的开源模块化单体；"
        "MySQL/Redis 保存业务事实，Outbox/Canal/Kafka 驱动派生数据更新。"
    )
    replace_labeled_paragraph(
        project_table.rows[1].cells[0].paragraphs[0],
        "项目背景：",
        background,
    )

    tech_stack = (
        "Java 21 / Spring Boot / Spring AI / MyBatis / MySQL / Redis / Kafka / Caffeine / Canal / Elasticsearch"
    )
    replace_labeled_paragraph(
        project_table.rows[2].cells[0].paragraphs[0],
        "技术栈：",
        tech_stack,
    )

    bullets = [
        (
            "高频互动计数：",
            "以分片 Bitmap 保存点赞/收藏状态，Lua 原子判重；Kafka 聚合增量并每秒折叠至定长 Redis SDS，缺失时基于"
            " BITCOUNT + Redisson 锁、限流和退避按需重建。",
        ),
        (
            "关系事件链路：",
            "将 following 作为主事实并在同一事务写 Outbox，经 Canal/Kafka 异步维护 follower、Redis ZSet 列表与用户"
            "计数，支持关系三态和时间游标分页。",
        ),
        (
            "Feed 缓存架构：",
            "拆分页面 ID 与内容片段，组合 Caffeine、Redis、MySQL 三级读取；以 SingleFlight、热点续期、TTL 抖动和反向索引"
            "控制击穿/雪崩，用户态在返回前实时叠加，避免污染公共缓存。",
        ),
        (
            "搜索与 RAG 问答：",
            "Elasticsearch 采用标题加权、标签过滤、function_score、search_after 与 Completion Suggester；RAG 以"
            " SHA-256/ETag 管理版本，按 800/100 字符滑窗切片，召回后按 postId 过滤并通过 SSE 流式回答。",
        ),
    ]
    detail_cell = project_table.rows[3].cells[0]
    while len(detail_cell.paragraphs) < len(bullets):
        detail_cell.add_paragraph()
    for paragraph, (label, body) in zip(detail_cell.paragraphs, bullets):
        replace_labeled_paragraph(paragraph, label, body)
    for paragraph in detail_cell.paragraphs[len(bullets):]:
        clear_paragraph(paragraph)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
