from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\Code\知光正版\zhiguang_be-main\output\简历1-项目1深度优化.docx")
OUTPUT = Path(r"D:\Code\知光正版\zhiguang_be-main\output\简历1-项目1产品包装版.docx")


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, east_asia, size_pt=10, bold=False, color="000000"):
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.color.rgb = None
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), east_asia)
    color_el = r_pr.find(qn("w:color"))
    if color_el is None:
        color_el = OxmlElement("w:color")
        r_pr.append(color_el)
    color_el.set(qn("w:val"), color)


def add_text(paragraph, text, *, east_asia="仿宋", size_pt=10, bold=False):
    run = paragraph.add_run(text)
    set_run_font(run, east_asia=east_asia, size_pt=size_pt, bold=bold)
    return run


def replace_labeled_paragraph(paragraph, label, body, *, size_pt=10):
    clear_paragraph(paragraph)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    add_text(paragraph, label, east_asia="黑体", size_pt=size_pt, bold=True)
    add_text(paragraph, body, east_asia="仿宋", size_pt=size_pt)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    project_table = document.tables[2]

    title_paragraph = project_table.rows[0].cells[0].paragraphs[0]
    clear_paragraph(title_paragraph)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(
        title_paragraph,
        "知光｜AI 增强型知识社区",
        east_asia="黑体",
        size_pt=10,
        bold=True,
    )

    context_cell = project_table.rows[1].cells[0]
    description = (
        "面向技术创作者与学习者，构建“创作发布—内容分发—互动沉淀—全文检索—语义问答”的产品闭环，"
        "支持 OSS 渐进发布、点赞/收藏、关注关系、个性化 Feed、ES 检索及单篇 RAG 问答。"
    )
    background = (
        "传统 CRUD 社区在高频互动、热点 Feed 回源、双向关系查询及搜索索引同步场景下，"
        "易出现写放大、缓存击穿和派生数据不一致；以 MySQL 为事实源、Redis 承载高频状态，"
        "通过 Outbox/Canal/Kafka 异步维护关系与搜索读模型。"
    )
    first_paragraph = context_cell.paragraphs[0]
    replace_labeled_paragraph(first_paragraph, "项目描述：", description, size_pt=9.5)
    if len(context_cell.paragraphs) < 2:
        second_paragraph = context_cell.add_paragraph()
    else:
        second_paragraph = context_cell.paragraphs[1]
    replace_labeled_paragraph(second_paragraph, "项目背景：", background, size_pt=9.5)
    for paragraph in context_cell.paragraphs[2:]:
        clear_paragraph(paragraph)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
