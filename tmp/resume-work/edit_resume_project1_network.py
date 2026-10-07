from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\Code\知光正版\zhiguang_be-main\output\简历1-项目1深度优化.docx")
OUTPUT = Path(r"D:\Code\知光正版\zhiguang_be-main\output\简历1-项目1网络包装版.docx")


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, east_asia, size_pt=10, bold=False, color="000000"):
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.color.rgb = None
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), east_asia)
    color_el = r_pr.find(qn("w:color"))
    if color_el is None:
        color_el = OxmlElement("w:color")
        r_pr.append(color_el)
    color_el.set(qn("w:val"), color)


def add_text(paragraph, text, *, east_asia="仿宋", size_pt=10, bold=False):
    run = paragraph.add_run(text)
    set_run_font(run, east_asia=east_asia, size_pt=size_pt, bold=bold)
    return run


def replace_labeled_paragraph(paragraph, label, body):
    clear_paragraph(paragraph)
    add_text(paragraph, label, east_asia="黑体", size_pt=10, bold=True)
    add_text(paragraph, body, east_asia="仿宋", size_pt=10)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    project_table = document.tables[2]

    title_paragraph = project_table.rows[0].cells[0].paragraphs[0]
    clear_paragraph(title_paragraph)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(
        title_paragraph,
        "知光｜知识社区与 RAG 问答平台",
        east_asia="黑体",
        size_pt=10,
        bold=True,
    )

    background = (
        "面向技术内容沉淀、社交互动与知识检索场景，构建覆盖内容发布、关注互动、个性化 Feed、全文检索和单篇文章 "
        "RAG 问答的开源模块化单体后端；针对高频互动计数、关系变更跨存储传播和索引更新问题，以 MySQL 为业务事实源，"
        "结合 Redis 与 Outbox/Canal/Kafka 维护派生数据。"
    )
    replace_labeled_paragraph(
        project_table.rows[1].cells[0].paragraphs[0],
        "项目背景：",
        background,
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
